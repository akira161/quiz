#!/usr/bin/env python3
"""過去問DB（専門・基礎・適性）をアプリ用に変換する。

使い方:
  python3 tools/build_db.py --senmon 専門.xlsx --kiso 基礎.xlsx --tekisei 適性.xlsx

入力: 「問題DB」シートと「問題画像」シートを持つ Excel（日本技術士会の過去問DB）
出力:
  data/senmon.xlsx, data/kiso.xlsx, data/tekisei.xlsx   画像を除いた軽量版（アプリが読み込む）
  figures/<科目>/<ID>.png                                問題の原本画像（階調を落として圧縮）
  precache.json                                          オフライン用に先に保存するファイル一覧

正答が空の行は tools/official_answers.json（日本技術士会の公式正答表）から補う。
専門科目は、従来の分類（3ブロック・サブジャンル・細目）と旧問題IDを tools/senmon_legacy.json から引き継ぐ。
必要なもの: openpyxl, pillow, numpy
"""
import argparse
import io
import json
import os
import re
import sys

import numpy as np
import openpyxl
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# 出典コード -> 年度表記（専門科目の従来ID「平成23年度_Ⅳ-1」と互換にするため従来の表記に合わせる）
def year_label(code):
    m = re.fullmatch(r'([HR])(\d+)(再?)', code)
    if not m:
        raise ValueError(code)
    era, n, retest = m.group(1), int(m.group(2)), m.group(3)
    if era == 'H':
        return f'平成{n}年度'
    if n == 1:
        return '令和元年度(再試験)' if retest else '令和元年度(本試験)'
    return f'令和{n}年度'

SUBJECTS = {
    # key: (出力ファイル名, 公式正答表の科目名, 題名に含まれる語)
    'senmon': ('senmon', None, '専門科目'),
    'kiso': ('kiso', '基礎', '基礎科目'),
    'tekisei': ('tekisei', '適性', '適性科目'),
}

# 図が必要そうな問題を文言から判定する（見落としても、アプリ側の「原本画像を見る」ボタンで確認できる）
FIGURE_NOUN = r'(?:図|表|グラフ|フロー|系統|回路|配置|構造|断面|平面|立面|概略|概念|模式|装置|ダイアグラム|曲線|線図|状態図|ネットワーク|プロセス|工程|設備|モデル)'
FIGURE_RE = re.compile(
    r'下図|下表|上図|上表|次図|次表|右図|左図|図中|表中|グラフ|ダイアグラム|フローシート|フロー図|系統図|回路図|断面図|配置図|模式図|概念図|線図'
    r'|図[0-9０-９一二三ＡＢＣABCa-cア-エ（(]|表[0-9０-９一二三ＡＢＣABCア-エ（(]'
    r'|[図表]に(?:示|記|あ)|[図表]の(?:よう|通り|とおり)|[図表]を(?:用|参|見)|[図表]から|[図表]は|[図表]の'
    r'|(?:以下|次|下|上)の[図表]|別[図表]|図示|[次下]に示[すし][。，、]'
    r'|(?:以下|次|下|上)に示す' + FIGURE_NOUN
)

def load_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    header_idx = next(i for i, r in enumerate(rows[:10]) if r and r[0] == 'ID')
    header = [str(c).strip() if c is not None else '' for c in rows[header_idx]]
    data = []
    for r in rows[header_idx + 1:]:
        if not r or not r[0]:
            continue
        data.append({h: v for h, v in zip(header, r) if h})
    return data

def extract_images(wb, ids_by_label):
    """「問題画像」シート内の画像を、画像の上端と同じ行にある「出典 問題番号」ラベルで問題IDに結びつける"""
    ws = wb['問題画像']
    labels = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v and str(v) in ids_by_label:
            labels[r] = str(v)
    images = {}
    for im in ws._images:
        row = im.anchor._from.row + 1
        label = labels.get(row)
        if label is None:
            raise RuntimeError(f'ラベルのない画像があります（行 {row}）')
        images[ids_by_label[label]] = im._data()
    return images

def compress(data):
    """スキャン画像の背景の汚れを白に飛ばし、階調を落として PNG-8 に圧縮する"""
    g = np.asarray(Image.open(io.BytesIO(data)).convert('L'))
    g = np.where(g >= 205, 255, g).astype(np.uint8)
    img = Image.fromarray(g, 'L').quantize(8, method=Image.Quantize.MEDIANCUT, dither=Image.Dither.NONE)
    buf = io.BytesIO()
    img.save(buf, 'PNG', optimize=True)
    return buf.getvalue()

def main():
    ap = argparse.ArgumentParser()
    for k in SUBJECTS:
        ap.add_argument(f'--{k}', required=True)
    args = ap.parse_args()

    official = json.load(open(os.path.join(ROOT, 'tools/official_answers.json'), encoding='utf-8'))['answers']
    legacy = json.load(open(os.path.join(ROOT, 'tools/senmon_legacy.json'), encoding='utf-8'))
    precache = []

    for key, (out_name, ans_subject, title_word) in SUBJECTS.items():
        path = getattr(args, key)
        wb = openpyxl.load_workbook(path)
        title = str(wb['問題DB']['A1'].value or '')
        if title_word not in title:
            sys.exit(f'{path}: 題名に「{title_word}」がありません（{title}）。科目の指定が違う可能性があります')
        rows = load_rows(wb['問題DB'])
        ids_by_label = {f"{r['出典']} {r['問題番号']}": r['ID'] for r in rows}
        images = extract_images(wb, ids_by_label)
        missing = [r['ID'] for r in rows if r['ID'] not in images]
        if missing:
            sys.exit(f'{key}: 画像のない問題があります: {missing[:5]}')

        # --- 正答 ---
        filled = 0
        for r in rows:
            a = str(r.get('正答') or '').strip()
            if not a:
                a = (official.get(f"{r['出典']}|{ans_subject}") or {}).get(r['問題番号'], '') if ans_subject else ''
                filled += 1 if a else 0
            if not re.fullmatch(r'[1-5](?:/[1-5])?|全員', a):
                sys.exit(f"{key}: 正答を決められません: {r['ID']} {r['問題番号']} -> {a!r}")
            r['正答'] = a

        # --- 専門科目: 従来の分類を引き継ぐ ---
        if key == 'senmon':
            if len(legacy) != len(rows):
                sys.exit('専門科目: 従来の分類データと問題数が合いません')
            for r, lg in zip(rows, legacy):
                if r['問題番号'] != lg['設問番号'] or year_label(r['出典']) != lg['年度']:
                    sys.exit(f"専門科目: 従来データと対応が合いません: {r['ID']}")
                r['旧問題ID'], r['ブロック'], r['旧サブジャンル'] = lg['旧問題ID'], lg['ブロック'], lg['旧サブジャンル']
                r['細目'], r['テーマ概要'], r['_legacy_fig'] = lg['細目'], lg['テーマ概要'], lg['図表']

        # --- 図が必要な問題 ---
        n_fig = 0
        for r in rows:
            text = str(r['問題文']) + ''.join(str(r.get(f'選択肢{c}') or '') for c in '①②③④⑤')
            need = bool(FIGURE_RE.search(text)) or bool(r.get('_legacy_fig'))
            r['図表'] = '要' if need else ''
            n_fig += need

        # --- 画像を保存 ---
        fig_dir = os.path.join(ROOT, 'figures', key)
        os.makedirs(fig_dir, exist_ok=True)
        for f in os.listdir(fig_dir):
            os.remove(os.path.join(fig_dir, f))
        total = 0
        for r in rows:
            png = compress(images[r['ID']])
            total += len(png)
            open(os.path.join(fig_dir, f"{r['ID']}.png"), 'wb').write(png)
            if r['図表'] == '要':
                precache.append(f"figures/{key}/{r['ID']}.png")

        # --- 軽量版 Excel ---
        cols = ['ID', '出典', '問題番号', 'Issue', '大分類', '小分類', '優先度', '正答', '問題文',
                '選択肢①', '選択肢②', '選択肢③', '選択肢④', '選択肢⑤', '図表']
        if key == 'senmon':
            cols += ['旧問題ID', 'ブロック', '旧サブジャンル', '細目', 'テーマ概要']
        out = openpyxl.Workbook()
        ws = out.active
        ws.title = '問題DB'
        ws.append(cols)
        for r in rows:
            ws.append([r.get(c) for c in cols])
        os.makedirs(os.path.join(ROOT, 'data'), exist_ok=True)
        out.save(os.path.join(ROOT, 'data', f'{out_name}.xlsx'))

        print(f'{key}: {len(rows)}問 / 正答補完 {filled}問 / 図表あり {n_fig}問 / 画像 {total / 1e6:.1f}MB')

    precache = sorted(precache)
    base = ['./', 'index.html', 'viewer.html', 'manifest.webmanifest', 'matome.xlsx',
            'data/senmon.xlsx', 'data/kiso.xlsx', 'data/tekisei.xlsx',
            'icons/icon-192.png', 'icons/icon-512.png', 'icons/apple-touch-icon.png']
    json.dump(base + precache, open(os.path.join(ROOT, 'precache.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
    print(f'precache.json: {len(base) + len(precache)}件')

if __name__ == '__main__':
    main()
